#!/usr/bin/env python
"""Ingest pharmacy product reviews (farmaline + medimarket xlsx) into the
mentions corpus, with rating-derived sentiment.

Each review row becomes a `Mention` (+ rating/brand_name/product_name) and a
`MentionClassification` whose sentiment is derived from the star rating:

    rating >= 4  -> positive
    rating == 3  -> neutral
    rating <= 2  -> negative

Topic/intent are left at their defaults here; `scripts/enrich_reviews_llm.py`
fills topic / side-effect / intent on the non-positive subset via the LLM.

Design notes
------------
* Streaming read (openpyxl read_only) — the farmaline file is ~435k rows / 92 MB.
* Bulk inserts in batches with ON CONFLICT DO NOTHING on text_hash (dedup),
  so the script is **idempotent / resumable** — re-running skips existing rows.
* Authors are pseudonymised (GDPR) before storage.
* No retention expiry is set — these are historical/public review records, so
  the 90-day raw-text wipe job must not touch them.
* Short reviews ("parfait") are kept — the connector's 30-char minimum is NOT
  applied; reviews are legitimately short. Empty bodies are skipped.

Usage:
    python scripts/ingest_reviews.py                # both files, full
    python scripts/ingest_reviews.py --limit 50000  # cap rows per file (demo)
    python scripts/ingest_reviews.py --only medimarket
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from sqlalchemy import create_engine, select, text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from core.config import settings
from core.security import pseudonymise_author
from ingestion.deduplication import sanitise_text
from models.data_source import DataSource, SourceType
from models.mention import Mention, MentionClassification, Sentiment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BATCH = 2000


# ── brand resolution ────────────────────────────────────────────────────────
def _load_brand_dictionary(engine) -> list:
    """Build a longest-first list of known brand prefixes.

    Seeded from medimarket's clean `brand` field (in raw_payload) once those
    rows exist, plus a curated multi-word set so farmaline's free-text
    productName (e.g. 'La Roche-Posay Effaclar …') resolves to a brand.
    """
    curated = [
        "La Roche-Posay", "Forte Pharma", "A.Vogel", "Be-Life", "Roger & Gallet",
        "CeraVe", "Eucerin", "Vichy", "Avène", "Bioderma", "Uriage", "Nuxe",
        "Mustela", "Puressentiel", "Pranarôm", "Metagenics", "Weleda", "Sanytol",
        "Hansaplast", "Tilman", "Louis Widmer", "Nutergia", "Aderma", "Ducray",
        "Klorane", "Phyto", "Caudalie", "Bach", "Boiron", "Arkopharma",
        "Solgar", "Bayer", "Sandoz", "Mylan", "EG", "Teva", "Nuxe", "Filorga",
        "ISDIN", "SVR", "Noreva", "Topicrem", "Cattier", "Sebamed", "Nivea",
    ]
    brands = set(curated)
    try:
        with engine.connect() as c:
            rows = c.execute(text(
                "select distinct brand_name from mentions "
                "where brand_name is not null and source_type='medimarket'"
            ))
            for (b,) in rows:
                if b:
                    brands.add(b)
    except Exception:
        pass
    # Longest first so 'La Roche-Posay' wins over 'La'.
    return sorted(brands, key=lambda s: -len(s))


def _match_brand(product_name: str, brand_prefixes: list) -> str | None:
    if not product_name:
        return None
    pn = product_name.strip()
    low = pn.lower()
    for b in brand_prefixes:
        bl = b.lower()
        if low == bl or low.startswith(bl + " "):
            return b
    # Fallback: first token (drops obvious noise tokens).
    first = pn.split()[0] if pn.split() else None
    if first and len(first) > 1 and first.lower() not in {"la", "le", "les", "de", "du"}:
        return first
    return None


def _sentiment_for(rating) -> Sentiment | None:
    if rating is None:
        return None
    try:
        r = float(rating)
    except (TypeError, ValueError):
        return None
    if r < 1:  # invalid (e.g. farmaline's -2)
        return None
    if r >= 4:
        return Sentiment.positive
    if r >= 3:
        return Sentiment.neutral
    return Sentiment.negative


def _parse_dt(val) -> datetime | None:
    if not val:
        return None
    if isinstance(val, datetime):
        return val if val.tzinfo else val.replace(tzinfo=timezone.utc)
    s = str(val).strip()
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None


def _text_hash(source_url: str | None, raw_text: str) -> str:
    import hashlib
    url_part = (source_url or "").strip()
    text_part = raw_text.strip()[:500]
    return hashlib.sha256(f"{url_part}||{text_part}".encode("utf-8")).hexdigest()


# ── source rows ──────────────────────────────────────────────────────────────
def _ensure_sources(session: Session) -> dict:
    specs = {
        "farmaline": dict(name="Farmaline (BE/FR pharmacy reviews)",
                          url="https://www.farmaline.be"),
        "medimarket": dict(name="Medi-Market (BE pharmacy reviews)",
                           url="https://www.medi-market.be"),
    }
    out = {}
    for key, spec in specs.items():
        ds = session.execute(
            select(DataSource).where(DataSource.name == spec["name"])
        ).scalar_one_or_none()
        if not ds:
            ds = DataSource(
                name=spec["name"],
                source_type=SourceType.pharmacy_import,
                url=spec["url"],
                country="BE",
                is_active=True,
                lawful_basis=settings.LAWFUL_BASIS,
                notes="Pharmacy product-review export imported via scripts/ingest_reviews.py",
            )
            session.add(ds)
            session.flush()
        out[key] = ds.id
    session.commit()
    return out


# ── row → mention dict mappers ────────────────────────────────────────────────
def _rows_medimarket(ws, source_id, limit):
    it = ws.iter_rows(min_row=2, values_only=True)
    n = 0
    for r in it:
        if limit and n >= limit:
            break
        n += 1
        review_id, source, product_id, product_url, author, age, rating, title, body, lang, rdate, raw_payload, scraped = r[:13]
        body = ("" if body is None else str(body)).strip()
        title = ("" if title is None else str(title)).strip()
        full = (f"{title}. {body}" if title and title not in body else body).strip()
        if not full:
            continue
        brand = None
        try:
            brand = (json.loads(raw_payload) or {}).get("brand") if raw_payload else None
        except Exception:
            brand = None
        yield {
            "source_id": source_id,
            "source_type": "medimarket",
            "source_url": product_url,
            "country": "BE",
            "language": (lang or "fr")[:5],
            "published_at": _parse_dt(rdate),
            "raw_text": full,
            "author_id_hash": pseudonymise_author(str(author), "medimarket") if author else None,
            "rating": int(float(rating)) if rating is not None else None,
            "brand_name": (brand or None),
            "product_name": None,
        }


def _rows_farmaline(ws, source_id, limit, brand_prefixes):
    it = ws.iter_rows(min_row=2, values_only=True)
    # cols: 0 title 1 upid 2 message 3 rating 5 productName 6 tenant 10 language
    #       13 submissionDate 16 customer.author
    n = 0
    tenant_country = {"be": "BE", "fr": "FR", "ch": "CH"}
    for r in it:
        if limit and n >= limit:
            break
        n += 1
        title = ("" if r[0] is None else str(r[0])).strip()
        msg = ("" if r[2] is None else str(r[2])).strip()
        rating = r[3]
        product_name = (("" if r[5] is None else str(r[5])).strip()) or None
        tenant = (r[6] or "be")
        lang = r[10]
        sub_date = r[13]
        author = r[16]
        full = (f"{title}. {msg}" if title and title not in msg else msg).strip()
        if not full:
            continue
        url = f"https://www.farmaline.be/p/{r[1]}" if r[1] else None
        yield {
            "source_id": source_id,
            "source_type": "farmaline",
            "source_url": url,
            "country": tenant_country.get(str(tenant).lower(), "BE"),
            "language": (lang or "fr")[:5],
            "published_at": _parse_dt(sub_date),
            "raw_text": full,
            "author_id_hash": pseudonymise_author(str(author), "farmaline") if author else None,
            "rating": int(float(rating)) if rating is not None else None,
            "brand_name": _match_brand(product_name, brand_prefixes),
            "product_name": product_name,
        }


# ── batch flush ───────────────────────────────────────────────────────────────
def _flush(engine, batch: list) -> int:
    """Insert a batch of mention dicts + their rating-derived classifications.

    Uses ON CONFLICT (text_hash) DO NOTHING and reads back which ids were
    actually inserted, then inserts matching classification rows.
    """
    if not batch:
        return 0
    now = datetime.now(timezone.utc)
    by_hash = {}  # dedupe within the batch — text_hash is unique
    sent_by_hash = {}
    for m in batch:
        th = _text_hash(m["source_url"], m["raw_text"])
        if th in by_hash:
            continue
        sent_by_hash[th] = _sentiment_for(m["rating"])
        by_hash[th] = {
            **m,
            "raw_text": m["raw_text"][:20000],
            "clean_text": sanitise_text(m["raw_text"])[:20000],
            "collected_at": now,
            "text_hash": th,
            "is_deleted": False,
            "retention_expires_at": None,  # historical reviews — never auto-wiped
        }
    mention_rows = list(by_hash.values())
    inserted = 0
    with Session(engine) as session:
        stmt = pg_insert(Mention).values(mention_rows)
        stmt = stmt.on_conflict_do_nothing(index_elements=["text_hash"])
        stmt = stmt.returning(Mention.id, Mention.text_hash)
        res = session.execute(stmt).all()
        cls_rows = []
        for mid, th in res:
            sent = sent_by_hash.get(th)
            cls_rows.append({
                "mention_id": mid,
                "sentiment": sent,
                "risk_type": "none",
                "is_adverse_event_candidate": False,
                "is_prescription_promotion": False,
                "model_name": "rating_heuristic",
                "review_status": "pending",
            })
        if cls_rows:
            cstmt = pg_insert(MentionClassification).values(cls_rows)
            cstmt = cstmt.on_conflict_do_nothing(index_elements=["mention_id"])
            session.execute(cstmt)
        inserted = len(res)
        session.commit()
    return inserted


FILES = {
    "medimarket": ("medimarket_reviews_20260521.xlsx", _rows_medimarket),
    "farmaline": ("farmaline_reviews.xlsx", _rows_farmaline),
}


def ingest_file(engine, key, limit, brand_prefixes, source_ids):
    fname, mapper = FILES[key]
    path = os.path.join(ROOT, fname)
    if not os.path.exists(path):
        print(f"[{key}] file not found: {path}", flush=True)
        return 0
    print(f"[{key}] opening {fname} …", flush=True)
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    src_id = source_ids[key]
    if key == "farmaline":
        gen = mapper(ws, src_id, limit, brand_prefixes)
    else:
        gen = mapper(ws, src_id, limit)
    batch, total, seen = [], 0, 0
    for row in gen:
        batch.append(row)
        seen += 1
        if len(batch) >= BATCH:
            total += _flush(engine, batch)
            batch = []
            if seen % 20000 == 0:
                print(f"[{key}] read {seen:,} | inserted {total:,}", flush=True)
    total += _flush(engine, batch)
    wb.close()
    print(f"[{key}] DONE read {seen:,} | inserted {total:,}", flush=True)
    return total


def _link_imported_reviews(engine) -> None:
    """Entity-link every review mention that names a tracked brand.

    CRITICAL: reviews are written with rating-derived classifications INLINE, which
    makes the async processing worker skip them (it only entity-resolves mentions
    that LACK a classification). Without this step the reviews sit in the corpus
    classified-but-unlinked, so the DIA engines / KPIs can't see them and render
    "Insufficient data" even though the data is right there. Running the dictionary
    resolver here closes the loop at import time. Idempotent (only mentions with no
    entity rows); deterministic; no LLM.
    """
    from processing.entity_resolution import get_dictionary
    from models.mention import MentionEntity, EntityType

    resolver = get_dictionary()
    if not getattr(resolver, "_loaded", False):
        resolver.load(settings.DATABASE_SYNC_URL)
    with Session(engine) as db:
        rows = db.execute(
            select(Mention.id, Mention.clean_text, Mention.language, Mention.country)
            .outerjoin(MentionEntity, MentionEntity.mention_id == Mention.id)
            .where(
                MentionEntity.id.is_(None),
                Mention.is_deleted == False,  # noqa: E712
                Mention.source_type.in_(("farmaline", "medimarket")),
                Mention.clean_text.isnot(None),
            )
        ).all()
        made = linked = 0
        for i, (mid, txt, lang, country) in enumerate(rows, start=1):
            ents = resolver.resolve_entities(txt or "", lang=lang or "en", country=country)
            seen = set()
            for e in ents:
                key = (e["entity_type"], e["entity_id"])
                if key in seen:
                    continue
                seen.add(key)
                db.add(MentionEntity(
                    mention_id=mid,
                    entity_type=EntityType(e["entity_type"]),
                    entity_id=int(e["entity_id"]),
                    confidence=round(float(e.get("confidence", 1.0)), 3),
                ))
                made += 1
            if ents:
                linked += 1
            if i % 1000 == 0:
                db.commit()
        db.commit()
        print(f"[link] entity-linked {linked}/{len(rows)} unlinked review mentions "
              f"({made} brand/product links)", flush=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0, help="max rows per file (0 = all)")
    ap.add_argument("--only", choices=list(FILES), default=None)
    args = ap.parse_args()

    engine = create_engine(settings.DATABASE_SYNC_URL)
    with Session(engine) as s:
        source_ids = _ensure_sources(s)
    # medimarket first so its clean brands seed the farmaline brand dictionary.
    order = [args.only] if args.only else ["medimarket", "farmaline"]
    grand = 0
    for key in order:
        prefixes = _load_brand_dictionary(engine) if key == "farmaline" else []
        grand += ingest_file(engine, key, args.limit or 0, prefixes, source_ids)
    print(f"TOTAL inserted this run: {grand:,}", flush=True)
    # Always close the loop: link the freshly-imported reviews to their brands so
    # the data actually feeds the KPIs (never leave them classified-but-unlinked).
    _link_imported_reviews(engine)


if __name__ == "__main__":
    main()
